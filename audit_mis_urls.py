#!/usr/bin/env python3
"""
audit_mis_urls.py
-----------------
Audits an Excel file of URLs for misplaced MIS (Industrial / Material Science)
product content. Outputs a new *_flagged.xlsx with classification columns appended.

Usage:
    python3 audit_mis_urls.py /path/to/file.xlsx
"""

import sys
import os
import re
import time
import argparse
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

REQUEST_TIMEOUT = 15          # seconds
REQUEST_DELAY   = 1.0         # seconds between requests
BLOCKED_CODES   = {403, 429}

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

# ---------------------------------------------------------------------------
# Keyword taxonomy
# Each entry: (category_label, [keywords], [model_number_patterns])
# ---------------------------------------------------------------------------

TAXONOMY = [
    (
        "Material Science / Industrial Microscopes",
        [
            "industrial microscope", "metallograph", "material science",
            "materials science", "hardness tester", "roughness",
            "stereo microscope", "zoom microscope", "measuring microscope",
            "toolmakers microscope", "toolmaker microscope",
            "industrial", "metrology", "surface measurement",
        ],
        [
            r"\bOLS\d*\b", r"\bLEXT\b", r"\bDSX\d*\b", r"\bBX53M\b",
            r"\bMX\d+\b", r"\bSTM\d*\b", r"\bSZX\d*\b", r"\bSZ\d+\b",
        ],
    ),
    (
        "Life Science",
        [
            "life science", "biological microscope", "fluorescence",
            "confocal", "cell imaging", "pathology",
        ],
        [
            r"\bBX\d+\b", r"\bCX\d+\b", r"\bIX\d+\b",
        ],
    ),
    (
        "Clinical",
        [
            "clinical", "hospital", "lab microscope",
            "hematology", "urinalysis",
        ],
        [
            r"\bCX43\b", r"\bCX33\b",
        ],
    ),
    (
        "Digital Cameras",
        [
            "digital camera", "microscope camera", "camera adapter",
        ],
        [
            r"\bDP\d+\b", r"\bSC\d+\b",
        ],
    ),
    (
        "Slide Scanners",
        [
            "slide scanner", "whole slide imaging", "digital pathology",
        ],
        [
            r"\bVS\d+\b",
        ],
    ),
    (
        "Objectives",
        [
            "objective lens", "UPlanSApo", "UIS2", "MPlanFL", "objective",
        ],
        [
            r"\bUPlanSApo\b", r"\bUIS2\b", r"\bMPlanFL\b",
        ],
    ),
]

# Flatten all model patterns for quick pre-check
ALL_MODEL_PATTERNS = [p for _, _, patterns in TAXONOMY for p in patterns]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def slug_text(url: str) -> str:
    """Return a human-readable version of the URL path for keyword matching."""
    parsed = urlparse(url)
    # Replace hyphens/underscores/slashes with spaces; lowercase
    raw = parsed.path + " " + parsed.query
    return re.sub(r"[-_/+]", " ", raw).lower()


def classify(text: str):
    """
    Run the taxonomy against *text* (lowercased).
    Returns (flagged: bool, categories: list, matched_kw: list, confidence: str)
    """
    text_lower = text.lower()
    matched_categories = []
    all_matched_kw = []
    total_hits = 0

    for category, keywords, model_patterns in TAXONOMY:
        cat_kw_hits = []
        cat_model_hits = []

        for kw in keywords:
            if kw.lower() in text_lower:
                cat_kw_hits.append(kw)

        for pattern in model_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            cat_model_hits.extend(matches)

        combined = cat_kw_hits + cat_model_hits
        if combined:
            matched_categories.append(category)
            all_matched_kw.extend(combined)
            total_hits += len(combined)

    flagged = len(matched_categories) > 0

    # Confidence
    has_model = any(
        re.search(p, text, re.IGNORECASE)
        for p in ALL_MODEL_PATTERNS
    )
    if total_hits >= 3 or has_model:
        confidence = "High"
    elif total_hits == 2:
        confidence = "Medium"
    elif total_hits == 1:
        confidence = "Low"
    else:
        confidence = ""

    return flagged, matched_categories, all_matched_kw, confidence


def scrape_lightweight(url: str):
    """
    Fetch only <title>, <meta description>, <h1> from a URL.
    Returns (text: str, method: str)
    method is one of: 'Live', 'Blocked', 'Error'
    """
    try:
        resp = requests.get(
            url,
            headers=HEADERS,
            timeout=REQUEST_TIMEOUT,
            allow_redirects=True,
        )
        if resp.status_code in BLOCKED_CODES:
            return "", "Blocked"

        soup = BeautifulSoup(resp.text, "lxml")

        parts = []

        title_tag = soup.find("title")
        if title_tag and title_tag.string:
            parts.append(title_tag.string.strip())

        meta_desc = soup.find("meta", attrs={"name": re.compile(r"description", re.I)})
        if meta_desc and meta_desc.get("content"):
            parts.append(meta_desc["content"].strip())

        for h1 in soup.find_all("h1"):
            parts.append(h1.get_text(" ", strip=True))

        return " ".join(parts), "Live"

    except requests.exceptions.Timeout:
        return "", "Blocked"
    except Exception:
        return "", "Error"


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def detect_url_column(ws):
    """
    Auto-detect the column index (1-based) that contains URLs.
    Scans up to first 20 rows, returns col index or None.
    """
    max_row_scan = min(20, ws.max_row or 20)
    for row in ws.iter_rows(min_row=1, max_row=max_row_scan):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and val.strip().lower().startswith("http"):
                return cell.column
    return None


def get_cell_value(ws, row, col):
    """Return cell value, resolving merged cells transparently."""
    cell = ws.cell(row=row, column=col)
    val = cell.value
    if val is None:
        # Check if it's part of a merged range
        for merged in ws.merged_cells.ranges:
            if (merged.min_row <= row <= merged.max_row and
                    merged.min_col <= col <= merged.max_col):
                # Value is on the top-left cell of the merged range
                val = ws.cell(row=merged.min_row, column=merged.min_col).value
                break
    return val


# ---------------------------------------------------------------------------
# Main audit logic
# ---------------------------------------------------------------------------

def audit(input_path: str):
    input_path = Path(input_path).resolve()
    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}")
        sys.exit(1)

    output_path = input_path.parent / (input_path.stem + "_flagged.xlsx")

    print(f"\n{'='*60}")
    print(f"  IMS / MIS URL Auditor")
    print(f"{'='*60}")
    print(f"  Input : {input_path}")
    print(f"  Output: {output_path}")
    print(f"{'='*60}\n")

    # Load workbook preserving all formatting
    wb = load_workbook(input_path)
    ws = wb.active

    # Detect URL column
    url_col = detect_url_column(ws)
    if url_col is None:
        print("ERROR: Could not detect a URL column (no cell starting with 'http' found in first 20 rows).")
        sys.exit(1)
    print(f"  Detected URL column: {openpyxl.utils.get_column_letter(url_col)} (col {url_col})\n")

    # Find the header row (first row with 'http' in the url column is data;
    # check if the row above is headers)
    header_row = None
    data_start_row = None
    for r in range(1, min(25, (ws.max_row or 25) + 1)):
        val = get_cell_value(ws, r, url_col)
        if isinstance(val, str) and val.strip().lower().startswith("http"):
            data_start_row = r
            if r > 1:
                header_row = r - 1
            break

    if data_start_row is None:
        print("ERROR: No URL rows found.")
        sys.exit(1)

    print(f"  Header row : {header_row}")
    print(f"  Data starts: row {data_start_row}\n")

    # Determine where to append new columns
    new_col_start = (ws.max_column or 1) + 1

    NEW_HEADERS = [
        "Flagged",
        "MIS Category",
        "Matched Keywords",
        "Confidence",
        "Scrape Method",
    ]

    # Write headers
    if header_row:
        for i, h in enumerate(NEW_HEADERS):
            cell = ws.cell(row=header_row, column=new_col_start + i, value=h)
            cell.font = Font(bold=True)

    # Counters
    total_urls    = 0
    total_flagged = 0
    blocked_count = 0
    category_counts = {}

    # Flag fill colours
    yes_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # light red
    no_fill  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # light green

    # Iterate data rows
    for row_idx in range(data_start_row, (ws.max_row or data_start_row) + 1):
        raw_val = get_cell_value(ws, row_idx, url_col)

        if not isinstance(raw_val, str):
            continue
        url = raw_val.strip()
        if not url.lower().startswith("http"):
            continue

        total_urls += 1
        print(f"  [{total_urls:>4}] {url[:80]}", end=" ... ", flush=True)

        # --- Attempt live scrape ---
        time.sleep(REQUEST_DELAY)
        live_text, scrape_method = scrape_lightweight(url)

        if scrape_method == "Blocked":
            blocked_count += 1
            # Fallback: slug only
            analysis_text = slug_text(url)
            scrape_method = "Slug-only"
            print("BLOCKED (slug fallback)", end=" ")
        elif scrape_method == "Error":
            analysis_text = slug_text(url)
            scrape_method = "Slug-only"
            print("ERROR (slug fallback)", end=" ")
        else:
            # Combine live text + slug for best coverage
            analysis_text = live_text + " " + slug_text(url)

        flagged, categories, matched_kw, confidence = classify(analysis_text)

        if flagged:
            total_flagged += 1
            for cat in categories:
                category_counts[cat] = category_counts.get(cat, 0) + 1
            print(f"FLAGGED [{confidence}] — {', '.join(categories)}")
        else:
            print("clean")

        # Write results
        flagged_str = "Yes" if flagged else "No"
        row_data = [
            flagged_str,
            "; ".join(categories) if categories else "",
            ", ".join(dict.fromkeys(matched_kw)),   # deduplicated, order-preserved
            confidence,
            scrape_method,
        ]
        for i, val in enumerate(row_data):
            cell = ws.cell(row=row_idx, column=new_col_start + i, value=val)
            if i == 0:  # Flagged column — colour code
                cell.fill = yes_fill if flagged else no_fill

    # Save
    wb.save(output_path)

    # Summary
    print(f"\n{'='*60}")
    print(f"  SUMMARY")
    print(f"{'='*60}")
    print(f"  Total URLs processed : {total_urls}")
    print(f"  Total flagged        : {total_flagged}")
    print(f"  Blocked / unreachable: {blocked_count}")
    if category_counts:
        print(f"\n  Flagged by category:")
        for cat, cnt in sorted(category_counts.items(), key=lambda x: -x[1]):
            print(f"    {cat:<45} {cnt}")
    print(f"\n  Output saved to:")
    print(f"    {output_path}")
    print(f"{'='*60}\n")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Audit an Excel file of URLs for misplaced MIS product content."
    )
    parser.add_argument(
        "xlsx_file",
        help="Path to the input .xlsx file (e.g. /path/to/Changes/file1.xlsx)",
    )
    args = parser.parse_args()
    audit(args.xlsx_file)
