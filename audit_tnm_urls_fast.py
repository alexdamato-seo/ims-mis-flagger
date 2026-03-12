#!/usr/bin/env python3
"""
audit_tnm_urls_fast.py
----------------------
Audits URLs for misplaced TNM/IMS (NDT / Testing / Inspection) content
that does NOT belong on a microscopy or life-science site.

Flags pages related to: Vanta, XRF, phased array, videoscopes, RVI,
IPLEX, 39DL Plus, Pipewizard, ultrasonic testing, remote visual inspection, etc.

Usage:
    python3 audit_tnm_urls_fast.py /path/to/file.xlsx
    python3 audit_tnm_urls_fast.py /path/to/file.xlsx --resume
"""

import sys
import re
import json
import argparse
from pathlib import Path
from urllib.parse import urlparse
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

REQUEST_TIMEOUT  = 15
MAX_WORKERS      = 10
CHECKPOINT_EVERY = 100
BLOCKED_CODES    = {403, 429}

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

# ---------------------------------------------------------------------------
# Keywords that clearly indicate a microscopy / life-science page.
# If the existing title matches any of these, skip full scraping.
# ---------------------------------------------------------------------------

MICROSCOPY_CLEAR_SIGNALS = [
    "microscope", "microscopy", "microscopi",
    "life science", "biological", "fluorescence", "confocal",
    "cell imaging", "pathology", "hematology", "urinalysis",
    "objective lens", "slide scanner", "whole slide",
    "digital pathology", "camera adapter", "microscope camera",
    "stereo microscope", "zoom microscope", "metallograph",
    "material science", "materials science",
]

# ---------------------------------------------------------------------------
# TNM / IMS keyword taxonomy
# Things that belong on an NDT/inspection site but NOT a microscopy/life-science site
# ---------------------------------------------------------------------------

TAXONOMY = [
    (
        "XRF / XRD Analyzers",
        [
            "xrf", "x-ray fluorescence", "xrd", "x-ray diffraction",
            "handheld xrf", "portable xrf", "pxrf", "alloy analysis",
            "elemental analysis", "metal verification", "positive material identification",
            "pmi",
        ],
        [
            r"\bVanta\b", r"\bDelta\b", r"\bOlympus\s*XRF\b",
            r"\bXpert\b", r"\bInnovX\b",
        ],
    ),
    (
        "Remote Visual Inspection (RVI)",
        [
            "remote visual inspection", "rvi", "videoscope", "video borescope",
            "borescope", "fiberscope", "industrial endoscope",
            "pipeline inspection", "turbine inspection",
        ],
        [
            r"\bIPLEX\b", r"\bIPlex\b",
        ],
    ),
    (
        "Ultrasonic Testing (UT)",
        [
            "ultrasonic testing", "ultrasonic inspection", "ultrasonic flaw",
            "phased array", "paut", "tofd", "time-of-flight diffraction",
            "thickness measurement", "wall thickness", "corrosion measurement",
            "weld inspection", "flaw detection", "pulse echo",
        ],
        [
            r"\b39DL\s*Plus\b", r"\b39DL\b", r"\bPipewizard\b",
            r"\bEpoch\b", r"\bOmniscan\b", r"\bFocusPX\b",
        ],
    ),
    (
        "Non-Destructive Testing (NDT/NDE)",
        [
            "non-destructive testing", "nondestructive testing", "ndt", "nde",
            "inspection system", "eddy current", "magnetic particle",
            "liquid penetrant", "acoustic emission",
        ],
        [
            r"\bNDT\b", r"\bNDE\b",
        ],
    ),
    (
        "Industrial Videoscopes / Inspection Tools",
        [
            "iplex", "3d measurement", "3d assist", "articulating probe",
            "insertion tube", "industrial camera inspection",
        ],
        [
            r"\bIPLEX\s*\w+\b", r"\b3DAssist\b",
        ],
    ),
]

ALL_MODEL_PATTERNS = [p for _, _, patterns in TAXONOMY for p in patterns]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def slug_text(url: str) -> str:
    parsed = urlparse(url)
    raw = parsed.path + " " + parsed.query
    return re.sub(r"[-_/+]", " ", raw).lower()


def classify(text: str):
    text_lower = text.lower()
    matched_categories = []
    all_matched_kw = []
    total_hits = 0

    for category, keywords, model_patterns in TAXONOMY:
        cat_kw_hits = []
        cat_model_hits = []

        for kw in keywords:
            if kw.lower() in text_lower:
                cat_kw_hits.append(kw)

        for pattern in model_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            cat_model_hits.extend(matches)

        combined = cat_kw_hits + cat_model_hits
        if combined:
            matched_categories.append(category)
            all_matched_kw.extend(combined)
            total_hits += len(combined)

    flagged = len(matched_categories) > 0

    has_model = any(
        re.search(p, text, re.IGNORECASE)
        for p in ALL_MODEL_PATTERNS
    )
    if total_hits >= 3 or has_model:
        confidence = "High"
    elif total_hits == 2:
        confidence = "Medium"
    elif total_hits == 1:
        confidence = "Low"
    else:
        confidence = ""

    return flagged, matched_categories, all_matched_kw, confidence


def scrape_content(url: str):
    """
    Fetch title, meta description, h1, h2, and up to 10 body paragraphs.
    Returns (text: str, method: str)
    """
    try:
        resp = requests.get(
            url,
            headers=HEADERS,
            timeout=REQUEST_TIMEOUT,
            allow_redirects=True,
        )
        if resp.status_code in BLOCKED_CODES:
            return "", "Blocked"

        soup = BeautifulSoup(resp.text, "lxml")
        parts = []

        title_tag = soup.find("title")
        if title_tag and title_tag.string:
            parts.append(title_tag.string.strip())

        meta_desc = soup.find("meta", attrs={"name": re.compile(r"description", re.I)})
        if meta_desc and meta_desc.get("content"):
            parts.append(meta_desc["content"].strip())

        for tag in soup.find_all(["h1", "h2"]):
            parts.append(tag.get_text(" ", strip=True))

        # Pull up to 10 body paragraphs for richer signal
        for p in soup.find_all("p")[:10]:
            text = p.get_text(" ", strip=True)
            if text:
                parts.append(text)

        return " ".join(parts), "Live"

    except requests.exceptions.Timeout:
        return "", "Blocked"
    except Exception:
        return "", "Error"


def is_clearly_microscopy(title: str) -> bool:
    """Return True if the title strongly suggests a microscopy/life-science page."""
    t = title.lower()
    return any(sig in t for sig in MICROSCOPY_CLEAR_SIGNALS)


def process_url(args):
    row_idx, url, existing_title = args

    # If the title clearly signals microscopy, skip the network fetch entirely
    if existing_title and is_clearly_microscopy(existing_title):
        title_text = existing_title + " " + slug_text(url)
        flagged, categories, matched_kw, confidence = classify(title_text)
        scrape_method = "Title-only (microscopy)"
    else:
        live_text, scrape_method = scrape_content(url)

        if scrape_method in ("Blocked", "Error"):
            analysis_text = (existing_title or "") + " " + slug_text(url)
            scrape_method = "Slug-only"
        else:
            analysis_text = live_text + " " + slug_text(url)

        flagged, categories, matched_kw, confidence = classify(analysis_text)

    return {
        "row_idx": row_idx,
        "url": url,
        "flagged": flagged,
        "categories": categories,
        "matched_kw": list(dict.fromkeys(matched_kw)),
        "confidence": confidence,
        "scrape_method": scrape_method,
    }


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def detect_url_column(ws):
    max_row_scan = min(20, ws.max_row or 20)
    for row in ws.iter_rows(min_row=1, max_row=max_row_scan):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and val.strip().lower().startswith("http"):
                return cell.column
    return None


def get_cell_value(ws, row, col):
    cell = ws.cell(row=row, column=col)
    val = cell.value
    if val is None:
        for merged in ws.merged_cells.ranges:
            if (merged.min_row <= row <= merged.max_row and
                    merged.min_col <= col <= merged.max_col):
                val = ws.cell(row=merged.min_row, column=merged.min_col).value
                break
    return val

# ---------------------------------------------------------------------------
# Main audit logic
# ---------------------------------------------------------------------------

def audit(input_path: str, resume: bool = False):
    input_path = Path(input_path).resolve()
    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}")
        sys.exit(1)

    output_path     = input_path.parent / (input_path.stem + "_tnm_flagged.xlsx")
    checkpoint_path = input_path.parent / (input_path.stem + "_tnm_checkpoint.json")

    print(f"\n{'='*60}")
    print(f"  TNM / IMS URL Auditor (FAST — {MAX_WORKERS} workers)")
    print(f"{'='*60}")
    print(f"  Input : {input_path}")
    print(f"  Output: {output_path}")
    print(f"{'='*60}\n")

    wb = load_workbook(input_path)
    ws = wb.active

    url_col = detect_url_column(ws)
    if url_col is None:
        print("ERROR: Could not detect a URL column.")
        sys.exit(1)
    print(f"  Detected URL column: {openpyxl.utils.get_column_letter(url_col)} (col {url_col})\n")

    header_row = None
    data_start_row = None
    for r in range(1, min(25, (ws.max_row or 25) + 1)):
        val = get_cell_value(ws, r, url_col)
        if isinstance(val, str) and val.strip().lower().startswith("http"):
            data_start_row = r
            if r > 1:
                header_row = r - 1
            break

    if data_start_row is None:
        print("ERROR: No URL rows found.")
        sys.exit(1)

    print(f"  Header row : {header_row}")
    print(f"  Data starts: row {data_start_row}\n")

    checkpoint = {}
    if resume and checkpoint_path.exists():
        with open(checkpoint_path) as f:
            checkpoint = json.load(f)
        print(f"  Resuming from checkpoint: {len(checkpoint)} rows already done\n")

    new_col_start = (ws.max_column or 1) + 1

    NEW_HEADERS = ["Flagged (TNM)", "TNM Category", "Matched Keywords", "Confidence", "Scrape Method"]

    if header_row:
        for i, h in enumerate(NEW_HEADERS):
            cell = ws.cell(row=header_row, column=new_col_start + i, value=h)
            cell.font = Font(bold=True)

    # Try to find a Page Title column (looks for "title" in header row)
    title_col = None
    if header_row:
        for c in range(1, (ws.max_column or 10) + 1):
            hval = str(ws.cell(row=header_row, column=c).value or "").lower()
            if "title" in hval:
                title_col = c
                break

    url_jobs = []
    for row_idx in range(data_start_row, (ws.max_row or data_start_row) + 1):
        raw_val = get_cell_value(ws, row_idx, url_col)
        if not isinstance(raw_val, str):
            continue
        url = raw_val.strip()
        if not url.lower().startswith("http"):
            continue
        if str(row_idx) in checkpoint:
            continue
        existing_title = ""
        if title_col:
            tv = get_cell_value(ws, row_idx, title_col)
            if isinstance(tv, str):
                existing_title = tv.strip()
        url_jobs.append((row_idx, url, existing_title))

    total_to_process = len(url_jobs) + len(checkpoint)
    print(f"  Total URLs   : {total_to_process}")
    print(f"  Already done : {len(checkpoint)}")
    print(f"  To fetch now : {len(url_jobs)}")
    print(f"  Workers      : {MAX_WORKERS}\n")

    yes_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    no_fill  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    results = dict(checkpoint)
    completed = 0

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_job = {executor.submit(process_url, job): job for job in url_jobs}

        for future in as_completed(future_to_job):
            result = future.result()
            row_idx = result["row_idx"]
            results[str(row_idx)] = result
            completed += 1

            status = (
                f"FLAGGED [{result['confidence']}] — {', '.join(result['categories'])}"
                if result["flagged"] else "clean"
            )
            print(f"  [{completed + len(checkpoint):>4}/{total_to_process}] {result['url'][:70]:<70} {status}")

            if completed % CHECKPOINT_EVERY == 0:
                with open(checkpoint_path, "w") as f:
                    json.dump(results, f)
                print(f"  >>> Checkpoint saved ({len(results)} rows)")

    print("\n  Writing results to Excel...")
    total_flagged = 0
    category_counts = {}

    for row_idx in range(data_start_row, (ws.max_row or data_start_row) + 1):
        key = str(row_idx)
        if key not in results:
            continue
        r = results[key]

        if r["flagged"]:
            total_flagged += 1
            for cat in r["categories"]:
                category_counts[cat] = category_counts.get(cat, 0) + 1

        row_data = [
            "Yes" if r["flagged"] else "No",
            "; ".join(r["categories"]) if r["categories"] else "",
            ", ".join(r["matched_kw"]),
            r["confidence"],
            r["scrape_method"],
        ]
        for i, val in enumerate(row_data):
            cell = ws.cell(row=row_idx, column=new_col_start + i, value=val)
            if i == 0:
                cell.fill = yes_fill if r["flagged"] else no_fill

    wb.save(output_path)

    if checkpoint_path.exists():
        checkpoint_path.unlink()

    total_urls = len(results)
    print(f"\n{'='*60}")
    print(f"  SUMMARY")
    print(f"{'='*60}")
    print(f"  Total URLs processed : {total_urls}")
    print(f"  Total flagged        : {total_flagged}")
    if category_counts:
        print(f"\n  Flagged by category:")
        for cat, cnt in sorted(category_counts.items(), key=lambda x: -x[1]):
            print(f"    {cat:<45} {cnt}")
    print(f"\n  Output saved to:")
    print(f"    {output_path}")
    print(f"{'='*60}\n")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Audit URLs for misplaced TNM/IMS (NDT/inspection) content."
    )
    parser.add_argument("xlsx_file", help="Path to the input .xlsx file")
    parser.add_argument("--resume", action="store_true", help="Resume from checkpoint if available")
    args = parser.parse_args()
    audit(args.xlsx_file, resume=args.resume)
