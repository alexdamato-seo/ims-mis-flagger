#!/usr/bin/env python3
"""
audit_downloads_ndt.py
----------------------
Scans MIS-DOWNLOADS.xlsx (or any downloads sheet) and flags rows
whose Page Title (col C) or URL slug indicate NDT / industrial inspection
content rather than IMS/microscopy content.

Classification is title-first, slug-assisted — no live fetching needed
since download titles are highly descriptive.

Usage:
    python3 audit_downloads_ndt.py MIS-DOWNLOADS.xlsx
"""

import sys
import re
from pathlib import Path

import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# ---------------------------------------------------------------------------
# NDT taxonomy
# ---------------------------------------------------------------------------

TAXONOMY = [
    (
        "Ultrasonic Testing (UT)",
        [
            "ultrasonic", "flaw detector", "thickness gauge",
            "thickness measurement", "pulse echo", "a-scan", "b-scan",
            "c-scan", "d-scan", "s-scan", "htha",
            "high temperature hydrogen attack",
            "corrosion mapping", "wall thickness", "pit depth",
            "immersion testing", "shear wave", "compression wave",
            "angle beam", "straight beam", "超声",
        ],
        [
            r"\bEPOCH\b", r"\b38DL\b", r"\b45MG\b", r"\b72DL\b",
            r"\bTG[- ]?\d+\b", r"\bUTWIN\b",
        ],
    ),
    (
        "Phased Array UT (PAUT)",
        [
            "phased array", "paut", "total focusing method", "tfm",
            "time of flight diffraction", "tofd",
            "full matrix capture", "fmc", "相控阵",
        ],
        [
            r"\bOmniScan\b", r"\bTomoScan\b", r"\bScanMaster\b",
            r"\bFOCUS LT\b", r"\bFOCUS\b",
        ],
    ),
    (
        "Eddy Current Testing (ECT)",
        [
            "eddy current", "eddy-current", "remote field",
            "bolthole", "bolt hole", "surface probe",
            "rotary probe", "pencil probe",
        ],
        [
            r"\bNORTEC\b", r"\bWeldSight\b",
        ],
    ),
    (
        "Visual / Remote Visual Inspection (RVI)",
        [
            "videoscope", "borescope", "video borescope",
            "endoscope", "industrial endoscope", "remote visual",
            "rigid borescope", "fiberscope", "visual inspection",
            "inspection kit", "太径硬性鏡", "細径硬性鏡",
            "目视检测", "硬性鏡", "视觉检测",
        ],
        [
            r"\bIPLEX\b", r"\bV-SCOPE\b", r"\bFB[- ]?[Ss]cope\b",
            r"\bRECON\b", r"\bSpitfire\b",
        ],
    ),
    (
        "XRF / Material Verification",
        [
            "xrf", "x-ray fluorescence", "material verification",
            "alloy verification", "alloy", "positive material identification",
            "pmi", "material analysis", "in-line automated",
        ],
        [
            r"\bVanta\b",
        ],
    ),
    (
        "NDE / NDT General",
        [
            "nde", "ndt", "non-destructive", "nondestructive",
            "inspection solution", "flexoform", "pipe inspection",
            "corrosion", "weld inspection", "缺陷", "检测解决方案",
            "open file format",
        ],
        [
            r"\bNDE\b", r"\bNDT\b",
        ],
    ),
]

ALL_MODEL_PATTERNS = [p for _, _, patterns in TAXONOMY for p in patterns]


# ---------------------------------------------------------------------------
# Classifier
# ---------------------------------------------------------------------------

def classify(text: str):
    text_lower = text.lower()
    matched_categories = []
    all_matched_kw = []
    total_hits = 0

    for category, keywords, model_patterns in TAXONOMY:
        cat_kw_hits = []
        cat_model_hits = []

        for kw in keywords:
            if kw.lower() in text_lower:
                cat_kw_hits.append(kw)

        for pattern in model_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            cat_model_hits.extend(matches)

        combined = cat_kw_hits + cat_model_hits
        if combined:
            matched_categories.append(category)
            all_matched_kw.extend(combined)
            total_hits += len(combined)

    flagged = len(matched_categories) > 0

    has_model = any(
        re.search(p, text, re.IGNORECASE)
        for p in ALL_MODEL_PATTERNS
    )
    if total_hits >= 3 or has_model:
        confidence = "High"
    elif total_hits == 2:
        confidence = "Medium"
    elif total_hits == 1:
        confidence = "Low"
    else:
        confidence = ""

    return flagged, matched_categories, list(dict.fromkeys(all_matched_kw)), confidence


def slug_text(url: str) -> str:
    """Convert URL path/query to space-separated words for keyword matching."""
    return re.sub(r"[-_/+?=\[\]0-9]", " ", url).lower()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def audit(input_path: str):
    input_path = Path(input_path).resolve()
    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}")
        sys.exit(1)

    output_path = input_path.parent / (input_path.stem + "_ndt_flagged.xlsx")

    print(f"\n{'='*60}")
    print(f"  MIS-DOWNLOADS NDT Auditor")
    print(f"{'='*60}")
    print(f"  Input : {input_path}")
    print(f"  Output: {output_path}")
    print(f"{'='*60}\n")

    wb = load_workbook(input_path)
    ws = wb.active

    # Detect columns from header row
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    title_col = headers.get("Page Title")
    url_col   = headers.get("Full Source URL") or headers.get("Source URL")

    if title_col is None:
        print("ERROR: 'Page Title' column not found.")
        sys.exit(1)

    print(f"  Page Title column : {openpyxl.utils.get_column_letter(title_col)}")
    if url_col:
        print(f"  URL column        : {openpyxl.utils.get_column_letter(url_col)}")

    new_col_start = ws.max_column + 1
    NEW_HEADERS = ["NDT Flagged", "NDT Category", "Matched Keywords", "Confidence"]
    for i, h in enumerate(NEW_HEADERS):
        cell = ws.cell(row=1, column=new_col_start + i, value=h)
        cell.font = Font(bold=True)

    yes_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    no_fill  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    total = 0
    flagged_count = 0
    category_counts = {}

    for row_idx in range(2, ws.max_row + 1):
        title = ws.cell(row_idx, title_col).value or ""
        url   = (ws.cell(row_idx, url_col).value or "") if url_col else ""

        analysis_text = title + " " + slug_text(url)
        flagged, categories, matched_kw, confidence = classify(analysis_text)

        total += 1
        if flagged:
            flagged_count += 1
            for cat in categories:
                category_counts[cat] = category_counts.get(cat, 0) + 1
            status = f"NDT [{confidence}] — {', '.join(categories)}"
        else:
            status = "IMS/clean"

        print(f"  [{row_idx-1:>4}] {str(title)[:55]:<55} {status}")

        row_data = [
            "Yes" if flagged else "No",
            "; ".join(categories) if categories else "",
            ", ".join(matched_kw),
            confidence,
        ]
        for i, val in enumerate(row_data):
            cell = ws.cell(row=row_idx, column=new_col_start + i, value=val)
            if i == 0:
                cell.fill = yes_fill if flagged else no_fill

    wb.save(output_path)

    print(f"\n{'='*60}")
    print(f"  SUMMARY")
    print(f"{'='*60}")
    print(f"  Total rows processed : {total}")
    print(f"  NDT flagged          : {flagged_count}")
    print(f"  IMS / clean          : {total - flagged_count}")
    if category_counts:
        print(f"\n  Flagged by NDT category:")
        for cat, cnt in sorted(category_counts.items(), key=lambda x: -x[1]):
            print(f"    {cat:<45} {cnt}")
    print(f"\n  Output saved to:")
    print(f"    {output_path}")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 audit_downloads_ndt.py <file.xlsx>")
        sys.exit(1)
    audit(sys.argv[1])
